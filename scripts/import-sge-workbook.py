"""Read the supplied workbook without modifying it. Output is private import data."""
import argparse, datetime, hashlib, json, pathlib
import openpyxl

def extract(workbook, estate_id):
    source = pathlib.Path(workbook)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    output = dict(version=1, estateId=estate_id, importId=digest, sourceFile=source.name,
                  blocks=[], harvesting=[], fieldActivities=[], removed=[])
    columns = {
        'Block Details': ('blocks', ['blockCode', 'plantingYear', 'plantingYearDescription', 'blockStatus', 'plantedHectares', 'plantedDate', 'plantingMaterial', 'soilType']),
        'DailyReception': ('harvesting', ['workDate', 'blockCode', 'employeeNo', 'employeeName', 'gang', 'activity', 'bunches']),
        'DailyActivity': ('fieldActivities', ['workDate', 'blockCode', 'gang', 'activityCode', 'activityDescription', 'mandays']),
    }
    expected_headers = {
        'Block Details': ['Block', 'Planting Year', 'Planting Year Description', 'BlockStatus', 'Planted Hactare', 'PlantedDate', 'PlantingMaterial', 'SoilType'],
        'DailyReception': ['Work Date', 'Block', 'Employee No', 'Employee Name', 'Gang', 'Activity', 'Bunches'],
        'DailyActivity': ['Work Date', 'Block', 'Gang', 'Activity Code', 'Activity Description', 'Mandays'],
    }
    for sheet, (table, names) in columns.items():
        ws = wb[sheet]
        headers = [v for v in next(ws.iter_rows(min_row=6, max_row=6, values_only=True))[:len(names)]]
        if headers != expected_headers[sheet]:
            raise ValueError(f'Unexpected headers in {sheet}: {headers}')
        seen = {}
        for row_no, values in enumerate(ws.iter_rows(min_row=7, values_only=True), 7):
            values = values[:len(names)]
            if all(v is None for v in values):
                continue
            values = [v.date().isoformat() if isinstance(v, datetime.datetime) else v for v in values]
            row = dict(zip(names, values))
            if not row['blockCode']:
                output['removed'].append(dict(sheet=sheet, row=row_no, reason='missing block'))
                continue
            key = json.dumps(values, ensure_ascii=False)
            if key in seen:
                output['removed'].append(dict(sheet=sheet, row=row_no, reason='exact duplicate', retainedRow=seen[key]))
                continue
            seen[key] = row_no
            for name in ['plantingYear', 'plantingYearDescription', 'employeeNo']:
                if name in row and row[name] is not None:
                    row[name] = str(row[name])
            if table != 'blocks':
                datetime.date.fromisoformat(row['workDate'])
                row['geolocation'] = None
            row.update(_id=f'{estate_id}-{table}-{digest[:12]}-{row_no}', estateId=estate_id,
                       sourceFile=source.name, sourceSheet=sheet, sourceRow=row_no, importId=digest)
            output[table].append(row)
    wb.close()
    output['summary'] = dict(blocks=len(output['blocks']), harvesting=len(output['harvesting']),
        fieldActivities=len(output['fieldActivities']), bunches=sum(r['bunches'] for r in output['harvesting']),
        mandays=sum(r['mandays'] for r in output['fieldActivities']), removed=len(output['removed']))
    return output

if __name__ == '__main__':
    p = argparse.ArgumentParser()
    p.add_argument('workbook'); p.add_argument('output'); p.add_argument('--estate', default='sg-gumut')
    args = p.parse_args()
    data = extract(args.workbook, args.estate)
    target = pathlib.Path(args.output)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    print(json.dumps(data['summary']))
