"""Extract only labelled COMMENT cells from monthly summaries, preserving annual source identity."""
import sys,json,hashlib
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
source,production,out=map(Path,sys.argv[1:4])
p=json.loads(production.read_text()); digest=hashlib.sha256(source.read_bytes()).hexdigest()
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
comments=[]; checks=[]; names={n['_id']:n for n in p['names']}
for year in range(2017,2027):
 sheet=f'Monthly {year}'; rows=list(w[sheet].values)
 records={r['sourceRow']:r for r in p['monthly'] if r['month']==f'{year}-01'}
 column=None; group=None
 for rn,row in enumerate(rows,1):
  if len(row)>1 and isinstance(row[1],str) and 'PERFORMANCE TABLE' in row[1]:group=row[1].replace(' PERFORMANCE TABLE','')
  headers=[i for i,v in enumerate(row) if isinstance(v,str) and v.strip().upper()=='COMMENT']
  if headers: column=headers[0]
  if rn not in records:continue
  r=records[rn]
  if column is not None:
   name=names[r['sourceBlockId']]
   assert str(row[2]).strip()==name['sourceBlockCode'] and group==name['division'],f'Block identity changed at {sheet}:{rn}'
  value=row[column] if column is not None and column<len(row) else None
  if isinstance(value,str) and value.strip() and not value.startswith('#'):
   key=[p['estateId'],r['sourceBlockId'],year]
   comments.append({'_id':'pc-'+hashlib.sha256(json.dumps(key,separators=(',',':')).encode()).hexdigest()[:40],
    'estateId':p['estateId'],'sourceBlockId':r['sourceBlockId'],'year':year,'comment':value.strip(),
    'sourceFile':source.name,'sourceSheet':sheet,'sourceRow':rn,'sourceCell':f'{get_column_letter(column+1)}{rn}',
    'sourceHash':digest,'sourceSystem':'oban-workbook','revision':0})
 # Independently verify all existing harvested-day values against the workbook cells.
 for r in [r for r in p['monthly'] if r['month'].startswith(str(year)) and r.get('harvestingDays') is not None]:
  ref=r['sourceCells']['harvestingDays'].split('!')[1]
  from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
  c,n=coordinate_from_string(ref)
  assert rows[n-1][column_index_from_string(c)-1]==r['harvestingDays'],ref
 checks.append({'year':year,'comments':sum(r['year']==year for r in comments)})
assert len({r['_id'] for r in comments})==len(comments)
out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps({'estateId':p['estateId'],'comments':comments,'checks':checks},separators=(',',':')))
print(json.dumps({'checks':checks,'comments':len(comments),'mapped':sum(bool(names[r['sourceBlockId']]['mapBlockCode']) for r in comments),'harvestingDaysVerified':sum(r.get('harvestingDays') is not None for r in p['monthly'])}))
