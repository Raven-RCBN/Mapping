"""Read the source workbook without editing it. Retain monthly rows and annual cells separately."""
import collections
import hashlib
import json
import pathlib
import sys
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

source = pathlib.Path(sys.argv[1])
destination = pathlib.Path(sys.argv[2])
workbook = openpyxl.load_workbook(source, data_only=True)
assert workbook.sheetnames[-1] == 'Wages 2025'
sha = hashlib.sha256(source.read_bytes()).hexdigest()
estate = 'f7b09538-fd0c-4ac3-b3ba-b6bf90e400ae'
minor = lambda value: int((Decimal(str(value)) * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
records = []
monthly_totals = collections.defaultdict(int)
monthly_groups = collections.defaultdict(int)
base = dict(currency='NGN', sourceSystem='oban-wages-2025-workbook', sourceFile=source.name,
            sourceHash=sha, sourceBlockCode=None, mapBlockCode=None)
for month, sheet in enumerate(workbook.worksheets[:-1], 1):
    total = 0
    for row in range(4, sheet.max_row + 1):
        division, serial, task, activity, amount = [sheet.cell(row, col).value for col in range(1, 6)]
        if not task or not activity:
            continue
        assert isinstance(amount, (int, float)), (sheet.title, row, 'Missing wage value')
        group = str(division).strip().upper() if division is not None else 'UNASSIGNED'
        value = minor(amount)
        record = dict(base, basis='monthly', period=f'2025-{month:02}', division=group,
                      task=str(task).strip(), activity=str(activity).strip(), amountMinor=value,
                      sourceDivision='' if division is None else str(division), sourceSerial=str(serial or ''),
                      sourceSheet=sheet.title, sourceRow=row, sourceCell=f'E{row}',
                      sourceKey=f'{sheet.title}!E{row}', note='Division-level wages. The source contains no block number.')
        records.append(record)
        monthly_groups[(group, record['task'], record['activity'])] += value
        monthly_totals[record['period']] += value
        total += value
    controls = [minor(sheet.cell(row, 5).value) for row in range(4, sheet.max_row + 1)
                if str(sheet.cell(row, 3).value).strip() == 'GRAND TOTAL:']
    assert controls == [total], (sheet.title, total, controls)
summary = workbook.worksheets[-1]
assert [summary.cell(3,c).value for c in [1,2,20]] == ['TASK NAME','ACTIVITY NAME','TOTAL WAGE']
annual_total = 0
for row in range(4, 161):
    task, activity = (str(summary.cell(row, col).value).strip() for col in (1,2))
    row_total = 0
    for col in range(3, 20):
        division = str(summary.cell(3, col).value)
        cell = summary.cell(row, col)
        assert isinstance(cell.value, (int, float)), cell.coordinate
        value = minor(cell.value)
        assert monthly_groups.pop((division, task, activity), 0) == value, cell.coordinate
        records.append(dict(base, basis='annual_summary', period='2025', division=division,
                            task=task, activity=activity, amountMinor=value, sourceDivision=division,
                            sourceSheet=summary.title, sourceRow=row, sourceCell=cell.coordinate,
                            sourceKey=f'{summary.title}!{cell.coordinate}',
                            note='Independent annual division summary. Do not add to monthly wages.'))
        annual_total += value
        row_total += value
    assert row_total == minor(summary.cell(row,20).value), row
assert not monthly_groups, monthly_groups
assert annual_total == minor(summary['T161'].value) == sum(monthly_totals.values())
for col in range(3,20):
    assert sum(minor(summary.cell(row,col).value) for row in range(4,161)) == minor(summary.cell(161,col).value)
production = json.loads(pathlib.Path('data/production/oban-production.json').read_text())
names = {n['_id']:n for n in production['names']}
groups = collections.defaultdict(set)
for record in production['yearly']:
    if record['year'] == 2025:
        name = names[record['sourceBlockId']]
        if name.get('mapBlockCode'):
            groups[name['mapBlockCode']].add(name['division'])
gis = json.loads(pathlib.Path('data/estates/oban-nigeria/gis-import.json').read_text())
polygons = {f['properties']['blockName'] for f in gis['boundary']['features']}
areas = []
for block in gis['blocks']:
    code = block['blockCode']
    if code not in polygons:
        continue
    group = block.get('division')
    note = 'Map GIS hectares, one allocation per polygon. Current map area used for 2025 estimates.'
    if not group and len(groups[code]) == 1:
        group = next(iter(groups[code]))
        note += ' Division confirmed by the 2025 production source.'
    if not group:
        group = 'DIVISION NOT SET'
        note += ' Set the division before this polygon participates in an allocation.'
    assert block['gisAreaHa'] > 0
    areas.append(dict(year=2025, sourceBlockId=block['_id'], sourceBlockCode=code,
                      division=group, mapBlockCode=code, mapHa=block['gisAreaHa'],
                      sourcePeriod=block.get('surveyDate',''), sourceFile=block['sourceFile'],
                      sourceSheet='GIS block inventory', sourceRow=block['sourceRow'],
                      sourceCell='gisAreaHa', note=note))
assert len(areas) == len({a['mapBlockCode'] for a in areas}) == 239
audit = dict(monthlyRecords=sum(r['basis']=='monthly' for r in records),
             annualSummaryRecords=sum(r['basis']=='annual_summary' for r in records),
             amountMinor=annual_total, currency='NGN', monthlyTotals=dict(monthly_totals),
             mappedRecords=0, blockNumbersPresent=False, annualCellsReconciled=2669,
             allocationAreas=len(areas), missingAreas=sum(a['mapHa'] is None for a in areas), unknownDivisions=sum(a['division']=='DIVISION NOT SET' for a in areas))
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(dict(version=1, estateId=estate, audit=audit, records=records, areas=areas), separators=(',',':')))
print(json.dumps(audit, indent=2))
