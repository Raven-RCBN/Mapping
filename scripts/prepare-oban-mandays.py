"""Extract annual activity mandays only; employee identities and source costs are never exported."""
import collections, hashlib, json, pathlib, sys
from decimal import Decimal
import openpyxl
source, wages_file, output = map(pathlib.Path, sys.argv[1:4])
wages = json.loads(wages_file.read_text())
keys = collections.defaultdict(set)
for r in wages['records']:
    if r['basis'] == 'monthly' and r['task'].upper() != 'MILL':
        keys[(r['division'], r['activity'].strip().upper())].add(r['task'])
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sums = collections.defaultdict(Decimal)
for sheet in book:
    if sheet.title not in {k[0] for k in keys}: continue
    headers = [c.value for c in next(sheet.iter_rows(min_row=3,max_row=3))]
    assert headers[2] == 'ActivityName' and headers[5] == 'Man-Day'
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row[0] or not row[2]: continue
        key = (sheet.title, str(row[2]).strip().upper())
        if key not in keys or len(keys[key]) != 1: continue
        assert isinstance(row[5], (float,int)) and row[5] >= 0
        sums[key] += Decimal(str(row[5]))
records = [dict(division=d, activity=a, task=next(iter(keys[(d,a)])), mandayUnits=int(v*1000000)) for (d,a),v in sorted(sums.items())]
payload = dict(version=1,estateId=wages['estateId'],year=2025,sourceSystem='oban-mandays-2025-workbook',sourceFile=source.name,sourceHash=hashlib.sha256(source.read_bytes()).hexdigest(),records=records)
output.write_text(json.dumps(payload,indent=2)+'\n')
print(json.dumps(dict(activityGroups=len(records),mandays=sum(r['mandayUnits'] for r in records)/1000000)))
