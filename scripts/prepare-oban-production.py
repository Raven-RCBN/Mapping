"""Read only monthly/yearly summaries; never open daily worksheet content or merge blocks."""
import sys,json,re,hashlib,collections
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as col,get_column_letter as letter
source,geo,out=map(Path,sys.argv[1:4])
w=openpyxl.load_workbook(source,read_only=True,data_only=True)
fw=openpyxl.load_workbook(source,read_only=True,data_only=False)
gis=json.loads(geo.read_text());estate=gis['estateId'];maps={b['blockCode']:b for b in gis['blocks']}
digest=hashlib.sha256(source.read_bytes()).hexdigest()
def norm(s):return re.sub(r'([A-Z])0+(\d)',r'\1\2',str(s).upper().strip())
def uid(prefix,*args):return prefix+hashlib.sha256(json.dumps(args,separators=(',',':')).encode()).hexdigest()[:40]
def number(v):return v if isinstance(v,(int,float)) and v>=0 else None
names={};monthly=[];yearly=[];checks=[]
def mapping(code,group):
 n=norm(code)
 if group=='2022':
  aliases={**{f'D{i}A':f'D-{i}' for i in range(1,6)},'E4':'E-4','E5':'E-5'}
  if n in aliases:return aliases[n],'confirmed_alias'
 if n in ['G1','G2'] and group in ['RO','YP']:return n+'A','confirmed_alias'
 shared={'F8':'F8+F9A','F9A':'F8+F9A','G8':'G8+G7A','G7A':'G8+G7A'}
 if n in shared:return shared[n],'shared_polygon'
 if n in maps:return n,'code'
 return None,'unmapped'
def getname(code,group,period):
 key=(code,group)
 if key not in names:
  mapped,method=mapping(code,group)
  names[key]={'_id':uid('pn-',estate,code,group),'estateId':estate,'sourceBlockCode':code,'division':group,
   'mapBlockCode':mapped,'linkMethod':method,'firstMonth':period,'lastMonth':period,
   'note':'Observed reporting periods, not an assumed renaming date. '+('Mapping confirmed by owner on 2026-10-10.' if method in ['confirmed_alias','shared_polygon'] else ''),'revision':0}
 names[key]['firstMonth']=min(period,names[key]['firstMonth']);names[key]['lastMonth']=max(period,names[key]['lastMonth'])
 return names[key]['_id']
def blockrows(rows,year):
 group='';header=[]
 for i,r in enumerate(rows,1):
  if r[1] and 'PERFORMANCE TABLE' in str(r[1]):group=str(r[1]).replace(' PERFORMANCE TABLE','')
  if year<=2019 and r[2] in ['OP MT','YP MT','RO MT']:group=str(r[2]).split()[0]
  if r[2]=='Blocks':header=r
  code=str(r[2]).strip() if r[2] is not None else ''
  if re.fullmatch(r'[A-Z]+\d+[AB]?|ESP',code) or (group=='DMC' and code in list('1234567')):yield i,code,group,r,header
def at(r,c):return r[col(c)-1] if c else None
for year in range(2017,2027):
 sheet=f'Monthly {year}';rows=list(w[sheet].values);forms=list(fw[sheet].values);count=0
 for rn,code,group,r,header in blockrows(rows,year):
  count+=1
  if year==2017:mtcols=[None]*3+['D']+[letter(i) for i in range(5,13)];bc='R';abw='AG';gps=eff=palms=None
  elif year==2018:mtcols=[letter(i) for i in range(4,16)];bc='W';abw='AM';gps='T';eff='S';palms=None
  elif year==2019:mtcols=[letter(i) for i in range(4,16)];bc='U';abw='AJ';gps=eff=palms=None
  elif year<=2021:mtcols=[letter(i) for i in range(37,49)];bc='V';abw='D' if year==2020 else 'G';gps='T' if year==2020 else 'E';eff='S' if year==2020 else 'F';palms=None
  elif year==2022:mtcols=[letter(i) for i in range(42,54)];bc='Y';abw='J';gps='G';eff='I';palms='F'
  else:mtcols=[letter(i) for i in range(58,70)];bc='AO';abw='J';gps='G';eff='I';palms='F'
  if year==2017 and group=='OP':mtcols[:4]=[None]*4
  for m in range(1,13):
   period=f'{year}-{m:02d}';sid=getname(code,group,period);issues=[];cells={}
   def read(field,c):
    v=at(r,c)
    if c:cells[field]=f'{sheet}!{c}{rn}'
    if isinstance(v,str) and v.startswith('#'):issues.append(f'{field}: {v} at {c}{rn}')
    return number(v)
   mt=read('mt',mtcols[m-1]);bunches=read('bunches',letter(col(bc)+m-1))
   status='not_received' if year==2026 and m>=8 else 'received' if mt is not None or bunches is not None else 'unavailable'
   effective=read('effectiveHa',eff);yieldarea=effective
   sourceyield=None
   if year>=2023:
    yc=letter(24+m);sourceyield=read('sourceYield',yc)
    formula=at(forms[rn-1],yc)
    match=re.search(r'/\$?([HI])\$?'+str(rn)+r'\b',str(formula))
    if match:yieldarea=read('yieldAreaHa',match.group(1))
   note=''
   if year==2017 and group=='OP' and m<=4:
    note=f'Only Jan–Apr combined tonnage supplied: {at(r,"D")} MT (D{rn}). Monthly tonnage unavailable; no allocation made.'
   dayscol=letter((col('BG') if year==2022 else col('BT'))+m-1) if year>=2022 else None
   days=read('harvestingDays',dayscol) if dayscol and header[col(dayscol)-1] in ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'] else None
   if group=='2022' and year>=2023:issues.append('Effective-area basis varies by source month; compare yield cautiously.')
   if mt and not yieldarea:issues.append('Effective area unavailable for reported production.')
   monthly.append({'_id':uid('mp-',estate,sid,period),'estateId':estate,'sourceBlockId':sid,'month':period,
    'receiptStatus':status,'mt':mt,'bunches':bunches,'workbookPalms':read('workbookPalms',palms),'gpsHa':read('gpsHa',gps),
    'effectiveHa':effective,'yieldAreaHa':yieldarea,'sourceAbw':read('sourceAbw',letter(col(abw)+m-1)),
    'sourceYield':sourceyield,'harvestingDays':days,
    'plantingYear':str(at(r,'D')) if year>=2022 and header[3]=='YOP' and number(at(r,'D')) else (group[:4] if re.match(r'^20\d\d',group) else ''),
    'parameterNote':'Parameters as supplied on this monthly summary; GIS survey counts retained separately.',
    'note':note,'sourceFile':source.name,'sourceSheet':sheet,'sourceRow':rn,'sourceSystem':'oban-workbook',
    'sourceKey':f'{sheet}:{rn}:{m}','sourceHash':digest,'sourceCells':cells,'issues':issues,'revision':0})
 checks.append({'sheet':sheet,'blocks':count,'monthlyRecords':count*12})
sheet='Yearly 2026';rows=list(w[sheet].values)
for rn,code,group,r,header in blockrows(rows,2026):
 for year in range(2017,2027):
  offset=2026-year;cols={'mt':76+offset,'bunches':62+offset,'workbookPalms':4+offset,'gpsHa':18+offset,
   'effectiveHa':29+offset,'sourceAbw':40+offset,'sourceBunchesPalm':51+offset,'sourceYield':87+offset}
  if all(r[i-1] is None for i in cols.values()):continue
  period=f'{year}-'+('07' if year==2026 else '12');sid=getname(code,group,period)
  values={k:number(r[i-1]) for k,i in cols.items()};issues=[f'{k}: {r[i-1]} at {letter(i)}{rn}' for k,i in cols.items() if isinstance(r[i-1],str) and r[i-1].startswith('#')]
  yearly.append({'_id':uid('yp-',estate,sid,year),'estateId':estate,'sourceBlockId':sid,'month':period,'year':year,
   'periodLabel':'Jan–Jul YTD' if year==2026 else 'Full year','receiptStatus':'received' if values['mt'] is not None or values['bunches'] is not None else 'unavailable',
   **values,'yieldAreaHa':values['effectiveHa'],'issues':issues,'sourceFile':source.name,'sourceSheet':sheet,'sourceRow':rn,
   'sourceSystem':'oban-workbook','sourceHash':digest,'sourceCells':{k:f'{sheet}!{letter(i)}{rn}' for k,i in cols.items()},'revision':0})
duplicates={k:[(r['sourceSheet'],r['sourceRow'],r['month']) for r in monthly if r['_id']==k] for k,n in collections.Counter(r['_id'] for r in monthly).items() if n>1}
assert not duplicates,json.dumps(duplicates)
assert len({r['_id'] for r in yearly})==len(yearly),'Duplicate annual identity'
latest=[r for r in monthly if r['month'][:4]=='2026']
assert abs(sum(r['mt'] or 0 for r in latest)-28871.27)<.001
nm={n['_id']:n for n in names.values()}
assert len([r for r in latest if r['month']=='2026-07' and nm[r['sourceBlockId']]['mapBlockCode']])==236
payload={'version':1,'estateId':estate,'sourceHash':digest,'sourceFile':source.name,'asOf':'2026-07-31',
 'names':list(names.values()),'monthly':monthly,'yearly':yearly,'checks':checks}
out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(payload,separators=(',',':')))
print(json.dumps({'names':len(names),'monthly':len(monthly),'yearly':len(yearly),'checks':checks,'2026mt':sum(r['mt'] or 0 for r in latest)},indent=2))
