"""Bounded, offline extraction. File contents are data, never instructions."""
import sys, json, csv, zipfile, datetime, subprocess, tempfile, os
from pathlib import Path
source, output, extension = sys.argv[1:4]
if sys.platform.startswith('linux'):
    import resource
    resource.setrlimit(resource.RLIMIT_AS,(2*1024**3,2*1024**3))
    resource.setrlimit(resource.RLIMIT_CPU,(360,360))
if len(sys.argv)>4 and sys.argv[4]=='render':
    page=int(sys.argv[5])
    if extension=='pdf':
        import pdfplumber
        with pdfplumber.open(source) as pdf:
            if page<1 or page>len(pdf.pages): raise ValueError('Invalid PDF page.')
        subprocess.run(['pdftoppm','-f',str(page),'-l',str(page),'-scale-to','5000','-png','-singlefile',source,output[:-4]],check=True,timeout=90,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
    else:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=100000000
        with Image.open(source) as image:
            image.seek(page-1)
            image.convert('RGB').save(output)
    print(json.dumps({'rendered':True}))
    sys.exit(0)
LIMIT=100000
count=0
bytes_written=0
warnings=[]
sheets=[]
def value(v):
    if isinstance(v, (datetime.datetime,datetime.date)): return v.isoformat()
    if v is None: return ''
    text=str(v)
    if len(text)>200000: raise ValueError('A source cell exceeds 200,000 characters. Split this file.')
    return text
with open(output,'w',encoding='utf8') as out:
    def emit(sheet, row, cells, kind='table'):
        global count, bytes_written
        count+=1
        if count>LIMIT: raise ValueError('More than 100,000 extracted rows. Split this file before importing.')
        if len(cells)>200: raise ValueError('More than 200 columns. Split the worksheet.')
        line=json.dumps(dict(sheet=sheet,sourceRow=row,kind=kind,cells=[value(v) for v in cells]),ensure_ascii=False)+'\n'
        bytes_written+=len(line.encode('utf8'))
        if bytes_written>512*1024*1024: raise ValueError('Extracted data exceeds 512 MB. Split this file.')
        out.write(line)
    def ocr(image):
        try:
            return subprocess.run(['tesseract',str(image),'stdout'],capture_output=True,timeout=60,check=True).stdout.decode('utf8')
        except FileNotFoundError: raise ValueError('OCR is unavailable on this server. Install Tesseract or upload a searchable PDF.')
    if extension in ['csv','tsv']:
        with open(source,encoding='utf-8-sig',newline='') as f:
            sample=f.read(65536); f.seek(0)
            try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
            except csv.Error: dialect=csv.excel_tab if extension=='tsv' else csv.excel
            for i,row in enumerate(csv.reader(f,dialect),1): emit('Data',i,row)
        sheets=['Data']
    elif extension=='xlsx':
        import openpyxl
        with zipfile.ZipFile(source) as z:
            if len(z.infolist())>10000 or sum(i.file_size for i in z.infolist())>500*1024*1024: raise ValueError('Workbook exceeds expanded size limit.')
        wb=openpyxl.load_workbook(source,read_only=True,data_only=False,keep_links=False)
        if len(wb.sheetnames)>50: raise ValueError('Maximum 50 worksheets per upload.')
        for ws in wb:
            sheets.append(ws.title)
            for i,row in enumerate(ws.iter_rows(values_only=True),1):
                if any(v is not None for v in row): emit(ws.title,i,row)
        warnings.append('Formulas are retained as source text, not executed. Review totals and formula cells before using them.')
        wb.close()
    elif extension=='xls':
        import xlrd
        wb=xlrd.open_workbook(source,on_demand=True)
        if wb.nsheets>50: raise ValueError('Maximum 50 worksheets per upload.')
        for ws in wb.sheets():
            sheets.append(ws.name)
            for i in range(ws.nrows): emit(ws.name,i+1,ws.row_values(i))
        warnings.append('Legacy Excel cached values were extracted; recalculate and export XLSX if totals need verification.')
    elif extension=='pdf':
        import pdfplumber
        with pdfplumber.open(source) as pdf:
            if len(pdf.pages)>200: raise ValueError('Maximum 200 PDF pages. Split this document.')
            for i,page in enumerate(pdf.pages,1):
                text=page.extract_text() or ''
                if not text.strip():
                    with tempfile.TemporaryDirectory() as tmp:
                        dest=str(Path(tmp)/'page')
                        subprocess.run(['pdftoppm','-f',str(i),'-l',str(i),'-scale-to','2400','-png','-singlefile',source,dest],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE,timeout=60)
                        text=ocr(dest+'.png')
                    warnings.append('Page %s used OCR; verify against the original.'%i)
                emit('Page '+str(i),i,[text], 'text')
                for t,table in enumerate(page.extract_tables(),1):
                    name='Page %s table %s'%(i,t); sheets.append(name)
                    for r,row in enumerate(table,1): emit(name,r,row)
        warnings.append('PDF table structure and OCR need review. Map drawings do not establish geographic coordinates.')
    elif extension in ['png','jpg','jpeg','webp','tif','tiff']:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=100000000
        with Image.open(source) as im:
            im.load()
            with tempfile.TemporaryDirectory() as tmp:
                image=Path(tmp)/'image.png'; im.convert('RGB').save(image)
                emit('Image OCR',1,[ocr(image)],'text')
        warnings.append('OCR text requires review. Use the image overlay option with geographic bounds to place this image on the map.')
    else: raise ValueError('Unsupported document format.')
print(json.dumps(dict(rows=count,sheets=sheets,warnings=warnings[:205])))
